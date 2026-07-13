from __future__ import annotations

import csv
import io
from datetime import date, datetime
from pathlib import Path
from typing import Any

from fastapi import APIRouter, Cookie, Depends, File, Form, HTTPException, Response, UploadFile, status
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy import or_
from sqlalchemy.orm import Session

from data.db_session import get_db
from data.pii import PIIDocument, PIIPerson
from data.users import User
from services.parsers import parse_file
from services.pii import storage as pii_storage
from services.pii.audit import log as pii_log
from services.pii.auth import (
    PII_TOKEN_TTL_SEC,
    issue_token,
    token_remaining_seconds,
    verify_token,
)
from services.pii.recognize import recognize_person
from utils.auth_deps import require_user
from utils.security import verify_password

router = APIRouter(prefix="/api/pii", tags=["pii"])

PII_COOKIE = "pii_token"


# ---------------------------------------------------------------------------
# Re-auth
# ---------------------------------------------------------------------------


class ReauthRequest(BaseModel):
    password: str


@router.post("/reauth")
async def pii_reauth(
    body: ReauthRequest,
    response: Response,
    user: User = Depends(require_user),
    db: Session = Depends(get_db),
):
    if not user.can_access_pii:
        pii_log(db, user.id, "reauth_fail", extra={"reason": "no_access"})
        raise HTTPException(status.HTTP_403_FORBIDDEN, "Доступ к персональным данным запрещён")
    if not verify_password(body.password, user.password_hash):
        pii_log(db, user.id, "reauth_fail", extra={"reason": "bad_password"})
        raise HTTPException(status.HTTP_401_UNAUTHORIZED, "Неверный пароль")

    token = issue_token(user.id)
    response.set_cookie(
        key=PII_COOKIE,
        value=token,
        max_age=PII_TOKEN_TTL_SEC,
        httponly=True,
        samesite="lax",
        path="/",
    )
    pii_log(db, user.id, "reauth_ok")
    return {"success": True, "expires_in": PII_TOKEN_TTL_SEC}


@router.post("/reauth/logout")
async def pii_reauth_logout(response: Response, user: User = Depends(require_user)):
    response.delete_cookie(PII_COOKIE, path="/")
    return {"success": True}


@router.get("/session")
async def pii_session_state(
    user: User = Depends(require_user),
    pii_token: str | None = Cookie(default=None),
):
    """Проверка состояния PII-доступа. Используется UI для решения, показывать ли модалку.
    remaining_seconds — реальный остаток до истечения токена (0 если просрочен/нет)."""
    can_access = bool(user.can_access_pii)
    remaining = token_remaining_seconds(pii_token, user.id) if can_access else 0
    return {
        "can_access": can_access,
        "active": remaining > 0,
        "remaining_seconds": remaining,
    }


# ---------------------------------------------------------------------------
# Dependency: проверка свежего PII-токена + аудит
# ---------------------------------------------------------------------------


def require_pii_access(
    user: User = Depends(require_user),
    pii_token: str | None = Cookie(default=None),
) -> User:
    if not user.can_access_pii:
        raise HTTPException(status.HTTP_403_FORBIDDEN, "Нет доступа к персональным данным")
    if not verify_token(pii_token, user.id):
        raise HTTPException(
            status.HTTP_401_UNAUTHORIZED,
            "Истёк токен доступа к персональным данным. Введите пароль заново.",
        )
    return user


# ---------------------------------------------------------------------------
# Persons
# ---------------------------------------------------------------------------


def _person_payload(p: PIIPerson, with_docs: bool = False, with_count: bool = True) -> dict[str, Any]:
    data: dict[str, Any] = {
        "id": p.id,
        "surname": p.surname,
        "name": p.name,
        "patronymic": p.patronymic,
        "birth_date": p.birth_date.isoformat() if p.birth_date else None,
        "full_name": p.full_name,
        "meta": p.meta or {},
        "created_at": p.created_at.isoformat(),
    }
    if with_count:
        data["documents_count"] = len(p.documents) if p.documents is not None else 0
    if with_docs:
        data["documents"] = [
            {
                "id": d.id,
                "filename": d.original_filename,
                "mime_type": d.mime_type,
                "size_bytes": d.size_bytes,
                "note": d.note,
                "uploaded_at": d.uploaded_at.isoformat(),
            }
            for d in sorted(p.documents, key=lambda x: x.uploaded_at, reverse=True)
        ]
    return data


@router.get("/persons")
async def list_persons(
    q: str | None = None,
    user: User = Depends(require_pii_access),
    db: Session = Depends(get_db),
):
    query = db.query(PIIPerson)
    if q:
        like = f"%{q.strip().lower()}%"
        query = query.filter(
            or_(
                PIIPerson.surname.ilike(like),
                PIIPerson.name.ilike(like),
                PIIPerson.patronymic.ilike(like),
            )
        )
    persons = query.order_by(PIIPerson.surname.asc(), PIIPerson.name.asc()).all()

    # Если по ФИО есть полные дубли — добавим в display подсказку с датой рождения
    items = [_person_payload(p) for p in persons]
    fios: dict[str, int] = {}
    for it in items:
        key = f"{it['surname']}|{it['name']}|{it['patronymic'] or ''}"
        fios[key] = fios.get(key, 0) + 1
    for it in items:
        key = f"{it['surname']}|{it['name']}|{it['patronymic'] or ''}"
        if fios[key] > 1 and it["birth_date"]:
            try:
                d = date.fromisoformat(it["birth_date"])
                it["full_name_with_dob"] = f"{it['full_name']} ({d.strftime('%d.%m.%Y')})"
            except ValueError:
                it["full_name_with_dob"] = it["full_name"]
        else:
            it["full_name_with_dob"] = it["full_name"]

    return {"success": True, "items": items}


@router.get("/persons/{person_id}")
async def get_person(
    person_id: int,
    user: User = Depends(require_pii_access),
    db: Session = Depends(get_db),
):
    p = db.get(PIIPerson, person_id)
    if not p:
        raise HTTPException(404, "Сотрудник не найден")
    pii_log(db, user.id, "view_person", entity="person", entity_id=p.id)
    return {"success": True, "person": _person_payload(p, with_docs=True)}


class PersonCreate(BaseModel):
    surname: str = Field(min_length=1, max_length=128)
    name: str = Field(min_length=1, max_length=128)
    patronymic: str | None = Field(default=None, max_length=128)
    birth_date: str | None = None  # YYYY-MM-DD


def _parse_birth_date(s: str | None) -> date | None:
    if not s:
        return None
    s = s.strip()
    if not s:
        return None
    # принимаем DD.MM.YYYY или YYYY-MM-DD
    for fmt in ("%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    raise HTTPException(400, "Неверный формат даты рождения")


@router.post("/persons")
async def create_person(
    body: PersonCreate,
    user: User = Depends(require_pii_access),
    db: Session = Depends(get_db),
):
    bd = _parse_birth_date(body.birth_date)
    existing = (
        db.query(PIIPerson)
        .filter(
            PIIPerson.surname == body.surname.strip(),
            PIIPerson.name == body.name.strip(),
            PIIPerson.patronymic == (body.patronymic.strip() if body.patronymic else None),
            PIIPerson.birth_date == bd,
        )
        .first()
    )
    if existing:
        return {"success": True, "person": _person_payload(existing), "created": False}

    p = PIIPerson(
        surname=body.surname.strip(),
        name=body.name.strip(),
        patronymic=(body.patronymic.strip() if body.patronymic else None),
        birth_date=bd,
    )
    db.add(p)
    db.commit()
    db.refresh(p)
    pii_log(db, user.id, "create_person", entity="person", entity_id=p.id)
    return {"success": True, "person": _person_payload(p), "created": True}


_PERSON_COL_ALIASES = {
    "surname": ("фамилия", "surname", "lastname", "last_name"),
    "name": ("имя", "name", "firstname", "first_name"),
    "patronymic": ("отчество", "patronymic", "middlename", "middle_name"),
    "birth_date": ("дата рождения", "датарождения", "дата_рождения", "birth_date", "birthdate", "birthday", "др"),
    "fullname": ("фио", "ф.и.о", "сотрудник", "fullname", "full_name", "физлицо", "физическое лицо"),
}


def _detect_person_columns(header: list) -> dict[str, int]:
    norm = [str(h or "").strip().lower() for h in header]
    mapping: dict[str, int] = {}
    for key, aliases in _PERSON_COL_ALIASES.items():
        for i, h in enumerate(norm):
            if h in aliases or any(a in h for a in aliases):
                mapping[key] = i
                break
    return mapping


def _rows_from_table(data: bytes, suffix: str) -> list[list[str]]:
    if suffix == ".csv":
        text = data.decode("utf-8-sig", errors="ignore")
        sample = text[:2000]
        delim = ";" if sample.count(";") >= sample.count(",") else ","
        return [list(r) for r in csv.reader(io.StringIO(text), delimiter=delim)]
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.active
    return [["" if c is None else str(c) for c in row] for row in ws.iter_rows(values_only=True)]


def _parse_person_table(data: bytes, suffix: str) -> list[dict]:
    # Старые форматы Excel (.xls/.ods) → конвертируем в .xlsx через LibreOffice (#9).
    if suffix in (".xls", ".ods"):
        import os
        import shutil as _sh
        import tempfile

        from services.parsers.office_convert import convert_to_modern

        fd, tmp = tempfile.mkstemp(suffix=suffix)
        with os.fdopen(fd, "wb") as f:
            f.write(data)
        try:
            conv = convert_to_modern(tmp)
            data = conv.read_bytes()
            _sh.rmtree(conv.parent, ignore_errors=True)
            suffix = ".xlsx"
        finally:
            try:
                os.unlink(tmp)
            except OSError:
                pass

    table = _rows_from_table(data, suffix)
    if not table:
        return []
    cols = _detect_person_columns(table[0])
    if not cols:
        return []
    has_header = any(k in cols for k in ("surname", "name", "fullname"))
    out: list[dict] = []
    for raw in table[1 if has_header else 0:]:
        def cell(key: str) -> str:
            i = cols.get(key)
            if i is None or i >= len(raw) or raw[i] is None:
                return ""
            return str(raw[i]).strip()

        surname, name = cell("surname"), cell("name")
        patr, bd = cell("patronymic"), cell("birth_date")
        if not surname and "fullname" in cols:
            parts = cell("fullname").split()
            if len(parts) >= 2:
                surname, name = parts[0], parts[1]
                patr = parts[2] if len(parts) >= 3 else ""
        if surname and name:
            out.append({
                "surname": surname, "name": name,
                "patronymic": patr or None, "birth_date": bd or None,
            })
    return out


@router.post("/import/1c")
async def import_1c_persons(
    file: UploadFile = File(...),
    user: User = Depends(require_pii_access),
    db: Session = Depends(get_db),
):
    """Импорт карточек сотрудников из табличной выгрузки 1С (CSV/XLSX/XLS/ODS).
    Колонки распознаются по заголовкам (Фамилия/Имя/Отчество/Дата рождения или ФИО).
    Персональные данные идут в раздел ПДн, НЕ в LLM/RAG (#18)."""
    suffix = Path(file.filename or "").suffix.lower()
    if suffix not in (".csv", ".xlsx", ".xls", ".ods"):
        raise HTTPException(400, "Ожидается таблица сотрудников из 1С (CSV/XLSX/XLS/ODS)")
    data = await _read_upload_file(file)
    try:
        persons = _parse_person_table(data, suffix)
    except Exception as e:
        raise HTTPException(400, f"Не удалось разобрать таблицу: {e}")
    if not persons:
        raise HTTPException(400, "Не найдены сотрудники. Нужны колонки «Фамилия»+«Имя» или «ФИО».")

    created = skipped = 0
    for r in persons:
        try:
            bd = _parse_birth_date(r.get("birth_date"))
        except HTTPException:
            bd = None
        existing = (
            db.query(PIIPerson)
            .filter(
                PIIPerson.surname == r["surname"],
                PIIPerson.name == r["name"],
                PIIPerson.patronymic == r["patronymic"],
                PIIPerson.birth_date == bd,
            )
            .first()
        )
        if existing:
            skipped += 1
            continue
        db.add(PIIPerson(surname=r["surname"], name=r["name"], patronymic=r["patronymic"], birth_date=bd))
        created += 1
    db.commit()
    pii_log(db, user.id, "import_1c", entity="person", extra={"created": created, "skipped": skipped})
    return {"success": True, "created": created, "skipped": skipped}


@router.delete("/persons/{person_id}")
async def delete_person(
    person_id: int,
    user: User = Depends(require_pii_access),
    db: Session = Depends(get_db),
):
    p = db.get(PIIPerson, person_id)
    if not p:
        raise HTTPException(404, "Сотрудник не найден")
    # Удаляем зашифрованные файлы с диска
    for d in p.documents:
        pii_storage.delete_file(d.storage_filename)
    db.delete(p)
    db.commit()
    pii_log(db, user.id, "delete_person", entity="person", entity_id=person_id)
    return {"success": True}


# ---------------------------------------------------------------------------
# Documents
# ---------------------------------------------------------------------------


_ALLOWED_PII_EXT = {".pdf", ".docx", ".doc", ".txt", ".md", ".jpg", ".jpeg", ".png"}
_MAX_PII_BYTES = 30 * 1024 * 1024  # 30 МБ


def _content_disposition(filename: str | None) -> str:
    """Безопасный Content-Disposition с кириллицей (ASCII-fallback + RFC 5987)."""
    from urllib.parse import quote

    name = filename or "document"
    ascii_name = name.encode("ascii", "ignore").decode().strip() or "document"
    return f"attachment; filename=\"{ascii_name}\"; filename*=UTF-8''{quote(name)}"


async def _read_upload_file(file: UploadFile) -> bytes:
    data = await file.read()
    if len(data) > _MAX_PII_BYTES:
        raise HTTPException(400, "Файл больше 30 МБ")
    return data


@router.post("/persons/{person_id}/documents")
async def upload_direct(
    person_id: int,
    file: UploadFile = File(...),
    note: str | None = Form(default=None),
    user: User = Depends(require_pii_access),
    db: Session = Depends(get_db),
):
    """Прямая загрузка документа в существующую группу. Без распознавания."""
    p = db.get(PIIPerson, person_id)
    if not p:
        raise HTTPException(404, "Сотрудник не найден")

    suffix = ("." + (file.filename or "").rsplit(".", 1)[-1].lower()) if "." in (file.filename or "") else ""
    if suffix not in _ALLOWED_PII_EXT:
        raise HTTPException(400, f"Неподдерживаемый формат: {suffix or '?'}")

    data = await _read_upload_file(file)
    storage_name, original_size = pii_storage.store_encrypted(data)

    doc = PIIDocument(
        person_id=p.id,
        original_filename=file.filename or storage_name,
        storage_filename=storage_name,
        mime_type=file.content_type,
        size_bytes=original_size,
        note=(note or None),
        uploaded_by=user.id,
    )
    db.add(doc)
    db.commit()
    db.refresh(doc)
    pii_log(db, user.id, "upload", entity="document", entity_id=doc.id, extra={"person_id": p.id})

    return {
        "success": True,
        "document": {
            "id": doc.id,
            "filename": doc.original_filename,
            "size_bytes": doc.size_bytes,
            "uploaded_at": doc.uploaded_at.isoformat(),
        },
    }


@router.post("/upload/quick-analyze")
async def upload_quick_analyze(
    file: UploadFile = File(...),
    user: User = Depends(require_pii_access),
    db: Session = Depends(get_db),
):
    """Быстрая загрузка: 1) парсим документ, 2) распознаём ФИО+дату рождения,
    3) возвращаем «кандидатов» — существующих людей с похожим ФИО.
    Файл при этом не сохраняем — фронт отдаст его ещё раз на /commit."""
    suffix = ("." + (file.filename or "").rsplit(".", 1)[-1].lower()) if "." in (file.filename or "") else ""
    if suffix not in (".pdf", ".docx", ".doc", ".txt", ".md", ".jpg", ".jpeg", ".png", ".webp", ".bmp", ".tif", ".tiff"):
        raise HTTPException(400, f"Формат «{suffix}» нельзя проанализировать. Используйте /persons/{{id}}/documents.")

    data = await _read_upload_file(file)

    # Распарсим текст
    import tempfile, os as _os
    fd, tmp_path = tempfile.mkstemp(suffix=suffix)
    try:
        with _os.fdopen(fd, "wb") as fp:
            fp.write(data)
        try:
            parsed = parse_file(tmp_path)
        except Exception as e:
            raise HTTPException(400, f"Не удалось распарсить файл: {e}")
    finally:
        try:
            _os.unlink(tmp_path)
        except OSError:
            pass

    recognized = recognize_person(parsed.text or "")

    # Кандидаты в БД
    candidates: list[dict] = []
    if recognized.get("surname"):
        q = db.query(PIIPerson).filter(PIIPerson.surname.ilike(recognized["surname"]))
        if recognized.get("name"):
            q = q.filter(PIIPerson.name.ilike(recognized["name"]))
        for p in q.limit(10).all():
            candidates.append(_person_payload(p))

    pii_log(db, user.id, "quick_analyze", extra={"filename": file.filename})

    return {
        "success": True,
        "filename": file.filename,
        "recognized": {
            "surname": recognized.get("surname"),
            "name": recognized.get("name"),
            "patronymic": recognized.get("patronymic"),
            "birth_date": recognized["birth_date"].isoformat() if recognized.get("birth_date") else None,
        },
        "candidates": candidates,
    }


@router.post("/upload/commit")
async def upload_commit(
    file: UploadFile = File(...),
    person_id: int | None = Form(default=None),
    surname: str | None = Form(default=None),
    name: str | None = Form(default=None),
    patronymic: str | None = Form(default=None),
    birth_date: str | None = Form(default=None),
    note: str | None = Form(default=None),
    user: User = Depends(require_pii_access),
    db: Session = Depends(get_db),
):
    """Окончательная загрузка после quick-analyze: либо привязываем к существующему person_id,
    либо создаём нового и привязываем."""
    suffix = ("." + (file.filename or "").rsplit(".", 1)[-1].lower()) if "." in (file.filename or "") else ""
    if suffix not in _ALLOWED_PII_EXT:
        raise HTTPException(400, f"Неподдерживаемый формат: {suffix or '?'}")

    data = await _read_upload_file(file)

    if person_id:
        person = db.get(PIIPerson, person_id)
        if not person:
            raise HTTPException(404, "Сотрудник не найден")
    else:
        if not (surname and name):
            raise HTTPException(400, "Не указан person_id и нет ФИО для создания группы")
        bd = _parse_birth_date(birth_date)
        # Проверка на существование (с учётом ДР)
        existing = (
            db.query(PIIPerson)
            .filter(
                PIIPerson.surname == surname.strip(),
                PIIPerson.name == name.strip(),
                PIIPerson.patronymic == (patronymic.strip() if patronymic else None),
                PIIPerson.birth_date == bd,
            )
            .first()
        )
        if existing:
            person = existing
        else:
            person = PIIPerson(
                surname=surname.strip(),
                name=name.strip(),
                patronymic=(patronymic.strip() if patronymic else None),
                birth_date=bd,
            )
            db.add(person)
            db.commit()
            db.refresh(person)
            pii_log(db, user.id, "create_person", entity="person", entity_id=person.id)

    storage_name, original_size = pii_storage.store_encrypted(data)
    doc = PIIDocument(
        person_id=person.id,
        original_filename=file.filename or storage_name,
        storage_filename=storage_name,
        mime_type=file.content_type,
        size_bytes=original_size,
        note=(note or None),
        uploaded_by=user.id,
    )
    db.add(doc)
    db.commit()
    db.refresh(doc)
    pii_log(db, user.id, "upload", entity="document", entity_id=doc.id, extra={"person_id": person.id})

    return {
        "success": True,
        "person_id": person.id,
        "document": {
            "id": doc.id,
            "filename": doc.original_filename,
            "size_bytes": doc.size_bytes,
            "uploaded_at": doc.uploaded_at.isoformat(),
        },
    }


@router.get("/documents/{document_id}/download")
async def download_pii_document(
    document_id: int,
    user: User = Depends(require_pii_access),
    db: Session = Depends(get_db),
):
    doc = db.get(PIIDocument, document_id)
    if not doc:
        raise HTTPException(404, "Документ не найден")
    try:
        data = pii_storage.load_decrypted(doc.storage_filename)
    except FileNotFoundError:
        raise HTTPException(404, "Файл отсутствует на диске")
    pii_log(db, user.id, "download", entity="document", entity_id=doc.id)

    return StreamingResponse(
        io.BytesIO(data),
        media_type=doc.mime_type or "application/octet-stream",
        headers={
            # RFC 5987: кириллические имена файлов нельзя класть в заголовок как есть
            # (latin-1) — иначе UnicodeEncodeError → 500. Даём ASCII-fallback + UTF-8.
            "Content-Disposition": _content_disposition(doc.original_filename),
            "Content-Length": str(len(data)),
        },
    )


@router.delete("/documents/{document_id}")
async def delete_pii_document(
    document_id: int,
    user: User = Depends(require_pii_access),
    db: Session = Depends(get_db),
):
    doc = db.get(PIIDocument, document_id)
    if not doc:
        raise HTTPException(404, "Документ не найден")
    person_id = doc.person_id
    pii_storage.delete_file(doc.storage_filename)
    db.delete(doc)
    db.commit()
    pii_log(db, user.id, "delete", entity="document", entity_id=document_id, extra={"person_id": person_id})
    return {"success": True}
